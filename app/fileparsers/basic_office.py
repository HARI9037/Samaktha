"""Bounded basic Office text extraction without model/OCR imports."""
from pathlib import Path

from app.fileparsers.base import DocumentParser, ParseResult


class BasicOfficeParser(DocumentParser):
    def can_handle(self, path: Path) -> bool:
        return path.suffix.lower() in {".docx", ".xlsx"}

    def parse(self, path: Path) -> ParseResult:
        try:
            if path.suffix.lower() == ".docx":
                from docx import Document
                document = Document(path)
                lines = [p.text for p in document.paragraphs]
                lines.extend("\t".join(c.text for c in row.cells) for table in document.tables for row in table.rows)
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    lines = []
                    for sheet in workbook:
                        lines.append(sheet.title)
                        for row in sheet.iter_rows():
                            lines.append("\t".join(str(c.value) if c.value is not None else "" for c in row))
                            if len(lines) >= 10000:
                                lines.append("[Extraction limited to 10000 rows]")
                                break
                        if len(lines) >= 10000:
                            break
                finally:
                    workbook.close()
            return ParseResult(ok=True, text="\n".join(lines), title=path.stem,
                               metadata={"parser": "basic_office", "format": path.suffix, "source_path": str(path)})
        except Exception:
            return ParseResult(ok=False, error="Basic Office text extraction failed; advanced document intelligence is disabled.")
