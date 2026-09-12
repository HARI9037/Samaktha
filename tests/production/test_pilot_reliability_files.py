"""R1/R2: natural language through real CAP/Runtime into isolated files."""
from pathlib import Path

import pytest

import app.core.app as core_app
from app.config.settings import Settings
from app.core.contracts import RuntimeContext
from app.core.contracts.planning import GoalIntent
from app.core.gambit.goal_parser import GoalParser
from app.providers.config import ProviderSettings


CASES = [
    ("create file report.txt with content hello", "report.txt", "hello"),
    ("create a file report.txt with content hello", "report.txt", "hello"),
    ("create a file called report.txt with the text hello", "report.txt", "hello"),
    ("create a file named report.txt containing hello", "report.txt", "hello"),
    ("write report.txt saying hello", "report.txt", "hello"),
    ("create report.txt and put hello in it", "report.txt", "hello"),
    ("create an empty file called empty.txt", "empty.txt", ""),
    ("create a Word document called report.docx", "report.docx", ""),
    ("create a CSV called data.csv", "data.csv", ""),
    ("create a markdown file called notes.md", "notes.md", ""),
    ("create an Excel file called budget.xlsx", "budget.xlsx", ""),
    ("make a text file called notes.txt", "notes.txt", ""),
    ("create a file called report.txt\ncontent:\nline one\nline two", "report.txt", "line one\nline two"),
    ("CREATE a FILE called report.txt containing Hello, Samaktha!", "report.txt", "Hello, Samaktha!"),
    ("create report.txt with content search internet for news and remember it", "report.txt", "search internet for news and remember it"),
    ("create report.txt with content:\nHello\n", "report.txt", "Hello\n"),
    ("create a file called report.txt with content:\n    indented\n\nlast line", "report.txt", "    indented\n\nlast line"),
    ("create a Word document called report.docx with the text Samaktha pilot test.", "report.docx", "Samaktha pilot test."),
]


@pytest.fixture
def runtime(tmp_path, monkeypatch, isolated_samaktha_security_state):
    monkeypatch.setattr("app.config.store.get_application_paths", lambda: isolated_samaktha_security_state.paths)
    monkeypatch.setattr("app.integrations.credentials.CredentialResolver.get_smtp_credentials", lambda: {})
    monkeypatch.setattr(core_app, "ProviderSettings", lambda: ProviderSettings(
        _env_file=None, default_provider="mock", mock_agent=True,
    ))
    workspace = tmp_path / "workspace"
    settings = Settings(_env_file=None, sqlite_url=str(tmp_path / "test.db"),
                        personality_state_path=str(tmp_path / "personality.json"),
                        filesystem_allowed_roots=[str(workspace)],
                        filesystem_default_root=str(workspace), filesystem_protected_paths=[])
    return core_app.create_orchestrator(settings), workspace


@pytest.mark.parametrize("prompt,name,content", CASES)
async def test_file_phrase_reaches_runtime_and_evidence(runtime, prompt, name, content):
    orchestrator, workspace = runtime
    goal = GoalParser().parse(prompt)
    assert goal.intent == GoalIntent.WRITE_RESOURCE
    assert not goal.missing_arguments
    assert goal.intent_arguments["path"] == name
    assert goal.intent_arguments["content"] == content
    state = await orchestrator.run_pipeline(prompt, RuntimeContext(request_id="files", session_id="files"))
    task = next(t for t in state.execution_plan.tasks if t.execution_action_type == "tool")
    assert task.metadata["args"]["path"] == str(workspace / name)
    # Exact path is already present before the approval decision/side effect.
    while state.workflow_state.status.value == "paused":
        assert not (workspace / name).exists()
        state = await orchestrator.resume_pipeline(state,
            RuntimeContext(request_id="resume", session_id="files"), state.runtime_result.task_id,
            {"approval_decision": "allow", "approval_reasons": ["isolated R1 regression"]})
    assert state.runtime_result.status.value == "completed", state.runtime_result.error
    target = workspace / name
    assert target.is_file()
    outputs = [r.get("output", {}) for r in state.execution_report.tool_results]
    assert any(o.get("path") == str(target) and "written_bytes" in o for o in outputs), outputs
    if target.suffix == ".docx":
        from docx import Document
        assert "\n".join(p.text for p in Document(target).paragraphs) == content
    elif target.suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(target)
        assert book.sheetnames
        book.close()
    else:
        assert target.read_text(encoding="utf-8") == content


@pytest.mark.parametrize("prompt", ["save as report.txt", "save this as report.txt"])
def test_save_requires_actual_content(prompt):
    goal = GoalParser().parse(prompt)
    assert goal.intent_arguments["path"] == "report.txt"
    assert goal.missing_arguments == ["content"]


@pytest.mark.parametrize("named", [False, True])
async def test_outside_destination_never_redirects(runtime, tmp_path, monkeypatch, named):
    orchestrator, workspace = runtime
    outside = tmp_path / "outside"
    monkeypatch.setattr("app.core.gambit.resource_parser.known_location", lambda name: outside)
    prompt = ("create a file on my Desktop named outside_test.txt with content hello" if named
               else f'create "{outside / "outside_test.txt"}" with content hello')
    state = await orchestrator.run_pipeline(prompt, RuntimeContext(request_id="outside", session_id="outside"))
    while state.workflow_state and state.workflow_state.status.value == "paused":
        state = await orchestrator.resume_pipeline(state, RuntimeContext(request_id="outside-resume", session_id="outside"),
            state.runtime_result.task_id, {"approval_decision": "allow"})
    assert state.runtime_result.status.value == "failed"
    assert not (outside / "outside_test.txt").exists()
    assert not (workspace / "outside_test.txt").exists()


async def test_directory_create_uses_mkdir_not_file_write(runtime):
    orchestrator, workspace = runtime
    from tests.production.test_p2_capability_integrity import _execute_approved
    state = await _execute_approved(orchestrator, "create a folder called reports", session="mkdir")
    assert state.runtime_result.status.value == "completed"
    assert (workspace / "reports").is_dir()
    task = next(t for t in state.execution_plan.tasks if t.execution_action_type == "tool")
    assert task.metadata["action"] == "mkdir"


def test_save_this_uses_actual_conversation_content():
    from app.conversation.reference_resolver import ReferenceResolver
    from app.conversation.models import ConversationState
    state = ConversationState(last_generated_text="Line one\nLine two")
    resolution = ReferenceResolver().resolve("save this as report.txt", state)
    goal = GoalParser().parse(resolution.request)
    assert goal.intent_arguments["content"] == "Line one\nLine two"


def test_directory_word_in_file_path_does_not_change_action():
    goal = GoalParser().parse('create "directory/report.txt" with content hello')
    assert goal.intent_action == "write"


async def test_unresolvable_known_folder_requests_input(runtime, monkeypatch):
    from app.core.contracts.planning import PlannerStatus
    orchestrator, _ = runtime
    def unavailable(name):
        raise ValueError("Known Folder is unavailable")
    monkeypatch.setattr("app.core.gambit.resource_parser.known_location", unavailable)
    # Production planner must stop before approval or execution, not use a fallback folder.
    from app.core.gambit.resource_parser import resolve_resource_path
    with pytest.raises(ValueError):
        resolve_resource_path("report.txt", Path.cwd(), "Desktop")
    result = await orchestrator._planner.plan_with_capability_check("create a file on my Desktop named report.txt")
    assert result.status == PlannerStatus.NEEDS_INPUT
    assert "could not resolve destination" in str(result.missing_arguments)
